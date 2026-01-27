from odoo import models, fields, api
from pypdf import PdfReader
from datetime import datetime
import base64
import re
from dateutil.relativedelta import relativedelta
import logging
import requests
from io import BytesIO
from pdf2image import convert_from_bytes
from PIL import Image
import pytesseract
import json
from docx import Document
from odoo.tools.mimetypes import guess_mimetype
_logger = logging.getLogger(__name__)


class UploadWidzard(models.Model):
    _name = 'upload.wizard'
    _description = 'For uploading the resume file'

    name = fields.Char("Name")
    json_data = fields.Text("Json Data")
    applicant_id = fields.Many2one('hr.applicant', string='Applicant')
    file_name = fields.Char('File Name')
    file_data = fields.Binary('File Data', required=True)
    skill_names = fields.Many2many('hr.applicant.skills', string='Skill Names')
    language_names = fields.Many2many('hr.applicant.language', string='Language Names')
    active = fields.Boolean(default=True)
    mimetype = fields.Char(
        compute="_compute_mimetype", string="Type", readonly=True, store=True
    )

    @api.depends("file_data")
    def _compute_mimetype(self):
        for record in self:
            binary = base64.b64decode(record.file_data or "")
            record.mimetype = guess_mimetype(binary)

    def action_ocr_resume(self, record=False):
        selected_record_ids = record
        if not record:
            selected_record_ids = self.browse(self.env.context.get('active_ids'))

        for selected_record_id in selected_record_ids:
            if not selected_record_id.file_data:
                selected_record_id.json_data = "No file content provided."
                continue

            binary = base64.b64decode(selected_record_id.file_data)
            file_type = selected_record_id.mimetype
            text_content = ""

            try:
                if file_type == "application/pdf":
                    try:
                        pdf_reader = PdfReader(BytesIO(binary))
                        for page in pdf_reader.pages:
                            page_text = page.extract_text()
                            text_content += page_text
                        if not text_content:
                            images = convert_from_bytes(binary, dpi=300)
                            extracted_text = []
                            for image in images:
                                text = pytesseract.image_to_string(image)
                                extracted_text.append(text)
                            text_content = "\n".join(extracted_text)

                    except Exception as e:
                        _logger.error("Error reading PDF: %s", str(e))
                        text_content = ""

                elif file_type in ["image/png", "image/jpeg", "image/jpg"]:
                    try:
                        image = Image.open(BytesIO(binary))
                        text_content = pytesseract.image_to_string(image)
                    except Exception as e:
                        _logger.error("Error processing image: %s", str(e))
                        text_content = ""

                elif file_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                    try:
                        doc = Document(BytesIO(binary))
                        text_content = ''

                        for section in doc.sections:
                            header = section.header
                            for paragraph in header.paragraphs:
                                text_content += paragraph.text + '\n'

                        for paragraph in doc.paragraphs:
                            text_content += paragraph.text + '\n'
                    except Exception as e:
                        _logger.error("Error processing DOCX: %s", str(e))
                        text_content = ""

                else:
                    _logger.error("Unsupported file type: %s", file_type)
                    selected_record_id.json_data = f"Unsupported file type: {file_type}"
                    continue

                json_response = selected_record_id.get_json_from_model(text_content)

                if json_response and "choices" in json_response and len(json_response["choices"]) > 0:
                    message_content = json_response["choices"][0].get("message", {}).get("content", "")

                    if message_content:
                        match = re.search(r'```json\n(.*?)\n```', message_content, re.DOTALL)
                        if match:
                            clean_json_str = match.group(1).strip()  # Extract JSON content
                            try:
                                parsed_json = json.loads(clean_json_str)
                                selected_record_id.json_data = json.dumps(parsed_json, indent=4)
                                file_name = parsed_json.get("name", "")
                                if file_name:
                                    self.write({'file_name': file_name})

                            except json.JSONDecodeError as e:
                                _logger.error("Error parsing JSON: %s", str(e))
                                selected_record_id.json_data = "Error parsing JSON"
                        else:
                            _logger.error("No valid JSON found in the content")
                            selected_record_id.json_data = "No valid JSON format found"
                    else:
                        _logger.error("No message content found in the response")
                        selected_record_id.json_data = "No message content found"
                else:
                    _logger.error("No valid response or choices in the API response")
                    selected_record_id.json_data = "No valid JSON data received."

            except Exception as e:
                _logger.error("Unexpected error during OCR processing: %s", str(e))
                selected_record_id.json_data = "An unexpected error occurred during file processing."

            _logger.info("Stored JSON data for file: %s", selected_record_id.file_name)


    def normalize_gender(self, gender):
        if gender:
            return gender.replace(" ", "").lower()
        return gender

    def normalize_marital_status(self, marital_status):
        if marital_status:
            return marital_status.replace(" ", "").lower()
        return marital_status

    def parse_experience(self, experience_str):
        years = 0
        months = 0

        year_match = re.search(r'(\d+)\s*[\+]*\s*years?', experience_str, re.IGNORECASE)
        month_match = re.search(r'(\d+)\s*months?', experience_str, re.IGNORECASE)

        if year_match:
            years = int(year_match.group(1))
        if month_match:
            months = int(month_match.group(1))

        return years, months

    def action_save(self):
        selected_record_ids = self.browse(self.env.context.get('active_ids'))
        created_applicants = []
        for file in selected_record_ids:
            if file.json_data:
                try:
                    parsed_json = json.loads(file.json_data)
                    partner_name = parsed_json.get("name", "")
                    email = parsed_json.get("email", "")
                    phone = parsed_json.get("phone", "")
                    degree_name = parsed_json.get("degree", "")
                    location = parsed_json.get("location", "")
                    experience = parsed_json.get("total_experience", "")
                    date_of_birth = parsed_json.get("date_of_birth", "")
                    marital_status = parsed_json.get("marital_status", "")
                    gender = parsed_json.get("gender", "")
                    file_data = file.file_data
                    linkedin = parsed_json.get("linkedin", "")
                    # Handle skills
                    skills_list = parsed_json.get("skills", [])
                    skills = []
                    if isinstance(skills_list, list):
                        for skill_name in skills_list:
                            existing_skill = self.env['hr.applicant.skills'].search([('name', '=', skill_name)],
                                                                                    limit=1)
                            if not existing_skill:
                                existing_skill = self.env['hr.applicant.skills'].create({'name': skill_name})
                            skills.append(existing_skill.id)
                    else:
                        _logger.warning("Skills are not in list format: %s", skills_list)

                    # Handle languages
                    languages_list = parsed_json.get("languages", [])
                    languages = []
                    if isinstance(languages_list, list):
                        for language_name in languages_list:
                            base_language = language_name.split("(")[
                                0].strip()  # Extract base language (e.g., "English" from "English (Fluent)")

                            # Search for an existing language ignoring proficiency level
                            existing_language = self.env['hr.applicant.language'].search(
                                [('name', '=ilike', base_language)], limit=1)

                            if not existing_language:
                                existing_language = self.env['hr.applicant.language'].create(
                                    {'name': base_language})

                            languages.append(existing_language.id)
                    else:
                        _logger.warning("Languages are not in list format: %s", languages_list)


                    gender = self.normalize_gender(gender)
                    valid_gender = dict(self.env['hr.applicant']._fields['gender'].selection).keys()
                    if gender not in valid_gender:
                        gender = 'not mentioned'

                    marital_status = self.normalize_marital_status(marital_status)
                    valid_marital_statuses = dict(self.env['hr.applicant']._fields['marital_status'].selection).keys()
                    if marital_status not in valid_marital_statuses:
                        marital_status = 'not mentioned'

                    if date_of_birth:
                        try:
                            date_of_birth = datetime.strptime(date_of_birth, '%d/%m/%Y').strftime('%Y-%m-%d')
                        except ValueError:
                            date_of_birth = None
                    else:
                        date_of_birth = None

                    experience = str(experience)
                    experience_years, experience_months = self.parse_experience(experience)
                    experience_years = str(experience_years)
                    experience_months = str(experience_months)

                    # new_candidate = selected_record_ids.env['hr.candidate'].create({
                    #     'partner_name': partner_name,
                    #     'email_from': email,
                    #     'partner_phone': phone,
                    # })

                    new_applicant = selected_record_ids.env['hr.applicant'].create({
                        'partner_name': partner_name,
                        'email_from': email,
                        'partner_phone': phone,
                        'degree': degree_name,
                        'location': location,
                        # 'experience_years': experience_years,
                        # 'experience_months': experience_months,
                        'skills': [(6, 0, skills)],
                        'languages': [(6, 0, languages)],
                        'date_of_birth': date_of_birth,
                        'marital_status': marital_status,
                        'gender': gender,
                        'linkedin_profile': linkedin,
                        'source_id': self.env.ref('sttl_recruitment_OCR.source_data_naukri').id,
                        # 'candidate_id': new_candidate.id,
                    })

                    attachment = self.env['ir.attachment'].create({
                        'name': partner_name,
                        'datas': file_data,
                        'res_model': 'hr.applicant',
                        'res_id': new_applicant.id,
                    })
                    new_applicant.write({'attachment_ids': [(4, attachment.id)]})
                    created_applicants.append(new_applicant.id)

                    file.write({'active': False})

                except json.JSONDecodeError as e:
                    _logger.error("Error parsing JSON data for file %s: %s", file.file_name, str(e))
            else:
                _logger.warning("No JSON data found for file %s", file.file_name)

        return {
            'name': "Created Applications",
            'type': 'ir.actions.act_window',
            'res_model': 'hr.applicant',
            'views': [[False, "list"], [False, "form"]],
            'view_mode': 'list,form',
            'domain': [('id',  'in',  created_applicants)],
            'target': 'current',
        }

    def get_json_from_model(self, text_content):
        api_url = "https://api.together.xyz/v1/chat/completions"
        qwen_api_key = self.env['ir.config_parameter'].sudo().get_param('sttl_recruitment_OCR.qwen_api_key')
        headers = {
            'Authorization': 'Bearer %s' % qwen_api_key,
            'Content-Type': 'application/json',
        }
        current_date = datetime.now()
        previous_month_date = current_date - relativedelta(months=1)
        previous_month_year = previous_month_date.strftime("%B %Y")
        payload = {
            "messages": [
                {
                    "role": "system",
                    "content": "provide the json data from the above content for below fields---\n\nname-- particularly the full name of the candidate\nskills-- the  skills of the candidate mentioned in the text specifically under skills section (add both soft and technical skills in this). \nemail -- the contact email of the candidate \nphone -- contact number of candidate usually a 10-12 number digits\ndegree-- the degree or qualification of the candidate mentioned in resume (only the name of the degrees or qualifications and not the whole details) "
                               f"\n experience in years and months in the format: Title of the experience (type of experience) (from month/year to  month/year) -> years and months (If the end date is marked as 'present' or 'till now', assume today's date is {previous_month_year} and calculate the months also properly). \n"
                               "\n Total Experience (Non-Overlapping) : in years and months \n"
                               "\n location-- search for the location or place mentioned in the resume where the candidate belong to. \n gender-- the gender of the candidate if mentioned. \ndate_of_birth-- the date of birth of the candidate in the format-%d/%m/%Y \nmarital_status-- the marital status of the candidate \n languages-- the languages known by the candidate mentioned in resume(not technical languages but the spoken ones specifically mentioned under languages section)\n"
                },
                {
                    "role": "user",
                    "content": text_content
                }
            ],
            "model": "Qwen/Qwen2.5-72B-Instruct-Turbo",
        }

        try:
            response = requests.post(api_url, json=payload, headers=headers)

            if response.status_code == 200:
                try:
                    json_response = response.json()
                    return json_response
                except ValueError as e:
                    _logger.error("Error parsing JSON: %s", str(e))
                    return {}
            else:
                _logger.error("Error in API call: %s", response.text)
                return {}

        except Exception as e:
            _logger.error("Exception during API call: %s", str(e))
            return {}